"""Tạo file mẫu import admin-ui/mau_import_ksk.xlsx (cần: pip install openpyxl).

Mọi ô dữ liệu được định dạng sẵn Text (@) để Excel KHÔNG tự đổi:
  00000001 -> 1,  10/10 -> 10-Oct,  01/01/1990 -> ngày kiểu Mỹ.
"""
from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.worksheet.datavalidation import DataValidation

HEADERS = [
    ("STT", 6), ("Họ và tên", 24), ("Mã nhân viên", 14), ("Giới tính", 9), ("Ngày sinh", 12),
    ("Chức danh", 22), ("Bộ phận", 16), ("Đơn vị", 20), ("Chiều cao", 9), ("Cân nặng", 9),
    ("Thể lực", 8), ("Huyết áp", 10), ("Mắt phải", 9), ("Mắt trái", 9), ("Bệnh mắt", 12),
    ("TMH", 8), ("RHM", 8), ("Nội da liễu", 10), ("Ngoại", 8), ("X quang", 9),
    ("Siêu âm", 20), ("Phụ khoa", 10), ("XN CTM", 8), ("XN nước tiểu", 11), ("Acid uric", 9),
    ("Creatinin", 9), ("Glucose", 8), ("Ure", 7), ("SGOT", 7), ("SGPT", 7),
    ("Triglycerides", 12), ("Cholesterol", 11), ("GGT", 7), ("Morphin", 10), ("Phân loại SK", 11),
    ("Kết luận", 30), ("Tư vấn", 36),
]
SAMPLE = [
    ["1", "NGUYỄN VĂN MẪU", "00000001", "Nam", "01/01/1990", "Nhân viên (dữ liệu giả)", "Bộ phận A", "Đơn vị Mẫu 1",
     "165", "60", "1", "120/80", "10/10", "10/10", "BT", "BT", "BT", "BT", "BT", "BT", "Bình thường", "",
     "BT", "BT", "320", "80", "5.0", "5.0", "20", "20", "1.5", "5.0", "25", "ÂM TÍNH", "Loại I",
     "Đủ sức khỏe làm việc (dữ liệu giả)", "Theo dõi sức khỏe định kỳ"],
    ["2", "TRẦN THỊ THỬ", "00000002", "Nữ", "15061995", "Nhân viên (dữ liệu giả)", "Bộ phận B", "Đơn vị Mẫu 2",
     "158", "50", "2", "110/70", "8/10", "9/10", "Cận nhẹ", "BT", "BT", "BT", "BT", "BT", "Bình thường", "BT",
     "BT", "BT", "280", "70", "4.8", "4.5", "18", "16", "1.2", "4.8", "18", "ÂM TÍNH", "Loại II",
     "Đủ sức khỏe làm việc (dữ liệu giả)", "Khám mắt, đo độ cận định kỳ"],
    ["3", "LÊ VĂN DEMO", "00000003", "Nam", "20/12/1985", "Nhân viên (dữ liệu giả)", "Bộ phận C", "Đơn vị Mẫu 1",
     "170", "72", "2", "130/85", "10/10", "10/10", "BT", "BT", "BT", "BT", "BT", "BT", "Gan nhiễm mỡ nhẹ (dữ liệu giả)", "",
     "BT", "BT", "420", "92", "5.8", "6.0", "30", "35", "2.4", "6.1", "40", "ÂM TÍNH", "Loại II",
     "Đủ sức khỏe làm việc (dữ liệu giả)", "Điều chỉnh chế độ ăn; tái khám sau 6 tháng"],
]
DATA_ROWS = 1000  # số dòng được định dạng sẵn Text

wb = Workbook()
ws = wb.active
ws.title = "Danh sách khám"
thin = Side(style="thin", color="CBD5E1")
head_fill = PatternFill("solid", fgColor="1E40AF")
req_fill = PatternFill("solid", fgColor="0D9488")
for c, (name, width) in enumerate(HEADERS, start=1):
    cell = ws.cell(row=1, column=c, value=name)
    required = name in ("Họ và tên", "Mã nhân viên", "Ngày sinh")
    cell.font = Font(bold=True, color="FFFFFF")
    cell.fill = req_fill if required else head_fill
    cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
    cell.border = Border(bottom=thin, right=thin)
    cell.number_format = "@"
    ws.column_dimensions[cell.column_letter].width = width
ws.row_dimensions[1].height = 32
for r in range(2, DATA_ROWS + 2):
    for c in range(1, len(HEADERS) + 1):
        ws.cell(row=r, column=c).number_format = "@"
for r, row in enumerate(SAMPLE, start=2):
    for c, v in enumerate(row, start=1):
        ws.cell(row=r, column=c, value=v)
ws.freeze_panes = "D2"
ws.auto_filter.ref = f"A1:{ws.cell(row=1, column=len(HEADERS)).column_letter}1"

dv = DataValidation(type="list", formula1='"Nam,Nữ"', allow_blank=True)
ws.add_data_validation(dv)
dv.add(f"D2:D{DATA_ROWS + 1}")

guide = wb.create_sheet("Hướng dẫn")
lines = [
    ("HƯỚNG DẪN NHẬP DANH SÁCH KHÁM SỨC KHỎE", True),
    ("", False),
    ("1. Nhập dữ liệu vào sheet \"Danh sách khám\". Xóa 3 dòng mẫu (dữ liệu giả) trước khi nhập thật.", False),
    ("2. Không đổi tên, không xóa cột ở dòng tiêu đề. Thứ tự cột có thể thay đổi.", False),
    ("3. Bắt buộc (tiêu đề màu xanh ngọc): Họ và tên, Mã nhân viên, Ngày sinh.", False),
    ("4. Ngày sinh: nhập 01/01/1990 hoặc 01011990 (ddmmyyyy).", False),
    ("5. Mọi ô đã định dạng Text, nên 00000001, 10/10, 110/70 được giữ nguyên.", False),
    ("   Khi dán dữ liệu từ file khác, dùng Dán đặc biệt > Giá trị (Ctrl+Alt+V > V) để không làm mất định dạng Text.", False),
    ("6. Ô để trống sẽ hiển thị \"—\" trên phiếu kết quả. Hệ thống không tự điền giá trị.", False),
    ("7. Lưu file dạng .xlsx rồi import tại trang Quản trị > Import dữ liệu > Kiểm tra file.", False),
    ("   Nếu có lỗi, hệ thống báo đúng số dòng và cột cần sửa, và chưa import gì cả.", False),
]
for i, (text, bold) in enumerate(lines, start=1):
    cell = guide.cell(row=i, column=1, value=text)
    cell.font = Font(bold=bold, size=13 if bold else 11)
guide.column_dimensions["A"].width = 110

out = Path(__file__).resolve().parent.parent / "admin-ui" / "mau_import_ksk.xlsx"
wb.save(out)
print("Đã tạo", out)
